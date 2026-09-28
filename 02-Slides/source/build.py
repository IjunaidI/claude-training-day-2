#!/usr/bin/env python3
"""Build the Day 2 deck: deck.html -> PDF (headless Chrome) -> PNGs -> PPTX with speaker notes.

    pip install -r requirements.txt
    python3 build.py              # writes ../ClaudeCode-Day2-IndusMotor.pdf and .pptx
    python3 build.py --pdf-only   # skip the PPTX

Chrome is found at the usual macOS/Linux/Windows paths, or set CHROME=/path/to/chrome.
Form links for the three QR codes live in config.json; edit them and rebuild.
"""
import html
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from html.parser import HTMLParser
from pathlib import Path

HERE = Path(__file__).resolve().parent
OUT_DIR = HERE.parent
NAME = "ClaudeCode-Day2-IndusMotor"

CHROME_CANDIDATES = [
    os.environ.get("CHROME", ""),
    "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
    "/usr/bin/google-chrome",
    "/usr/bin/chromium",
    r"C:\Program Files\Google\Chrome\Application\chrome.exe",
]


def chrome() -> str:
    for c in CHROME_CANDIDATES:
        if c and Path(c).exists():
            return c
    sys.exit("Chrome not found. Set CHROME=/path/to/chrome.")


def make_qr_codes() -> None:
    import qrcode
    from qrcode.image.pil import PilImage

    cfg = json.loads((HERE / "config.json").read_text())
    for key in ("quiz_a", "quiz_b", "feedback"):
        img = qrcode.make(cfg[key]["url"], image_factory=PilImage, border=0, box_size=20)
        img.save(HERE / "assets" / f"qr-{key}.png")
    # the link text under each QR comes from config.json too
    deck = (HERE / "deck.html").read_text()
    for key in ("quiz_a", "quiz_b", "feedback"):
        deck = re.sub(
            rf'(<a class="qlink" data-key="{key}"[^>]*>)(.*?)(</a>)',
            lambda m: m.group(1) + html.escape(cfg[key]["url"]) + m.group(3),
            deck,
        )
    (HERE / "deck.html").write_text(deck)


def render_pdf(pdf: Path) -> None:
    subprocess.run(
        [chrome(), "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
         "--virtual-time-budget=4000", f"--print-to-pdf={pdf}", (HERE / "deck.html").as_uri()],
        check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
    )


class NotesParser(HTMLParser):
    """Collects the text of every <aside class="notes"> in slide order."""

    def __init__(self):
        super().__init__()
        self.notes, self._depth, self._buf, self._slides = [], 0, [], 0

    def handle_starttag(self, tag, attrs):
        cls = dict(attrs).get("class", "") or ""
        if tag == "section" and "slide" in cls.split():
            self._slides += 1
            self.notes.append("")
        if tag == "aside" and "notes" in cls.split():
            self._depth, self._buf = 1, []
        elif self._depth and tag in ("p", "br", "li"):
            self._buf.append("\n")

    def handle_endtag(self, tag):
        if tag == "aside" and self._depth:
            text = re.sub(r"[ \t]+", " ", "".join(self._buf))
            text = "\n".join(line.strip() for line in text.splitlines())
            self.notes[-1] = re.sub(r"\n{3,}", "\n\n", text).strip()
            self._depth = 0

    def handle_data(self, data):
        if self._depth:
            self._buf.append(data)


def render_pptx(pdf: Path, pptx: Path) -> None:
    from pptx import Presentation
    from pptx.util import Emu

    parser = NotesParser()
    parser.feed((HERE / "deck.html").read_text())

    tmp = Path(tempfile.mkdtemp())
    subprocess.run(["pdftoppm", "-png", "-scale-to-x", "1920", "-scale-to-y", "1080", str(pdf), str(tmp / "s")], check=True)
    pngs = sorted(tmp.glob("s-*.png"))
    if len(pngs) != len(parser.notes):
        print(f"warning: {len(pngs)} pages but {len(parser.notes)} slides with notes")

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    blank = prs.slide_layouts[6]
    for i, png in enumerate(pngs):
        s = prs.slides.add_slide(blank)
        s.shapes.add_picture(str(png), 0, 0, prs.slide_width, prs.slide_height)
        if i < len(parser.notes) and parser.notes[i]:
            s.notes_slide.notes_text_frame.text = parser.notes[i]
    prs.save(pptx)
    shutil.rmtree(tmp)


def main() -> None:
    make_qr_codes()
    pdf = OUT_DIR / f"{NAME}.pdf"
    render_pdf(pdf)
    print(f"wrote {pdf}")
    if "--pdf-only" not in sys.argv:
        pptx = OUT_DIR / f"{NAME}.pptx"
        render_pptx(pdf, pptx)
        print(f"wrote {pptx}")


if __name__ == "__main__":
    main()
